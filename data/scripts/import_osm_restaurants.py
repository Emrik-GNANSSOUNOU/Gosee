"""
Importe des restaurants réels depuis OpenStreetMap (API Overpass) dans
benin_contenu_curation.xlsx, feuille « Restaurants (OSM) », pour les pauses
repas de « Créer sa sortie ».

Sélection : restaurants nommés situés à moins de RAYON_KM d'un lieu du
catalogue (site, loisir, activité) — inutile d'importer des restaurants loin
de toute sortie possible —, au plus MAX_PAR_LIEU par lieu, en privilégiant
les fiches OSM les plus complètes (horaires, téléphone, cuisine, site web).
Aucune information n'est inventée : description générée uniquement à partir
des tags OSM et de la distance au lieu voisin, prix laissés vides (OSM ne les
renseigne pas).

Licence : données © contributeurs OpenStreetMap, sous licence ODbL —
l'attribution est affichée sur les fiches (colonne Statut = « OpenStreetMap »).

Réseau requis : overpass-api.de.

Usage : python data/scripts/import_osm_restaurants.py
puis    python data/scripts/export_xlsx_to_json.py
        python data/scripts/generate_sql_update.py
"""
import json
import math
import re
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX_PATH = ROOT / "benin_contenu_curation.xlsx"
LIEUX_JSON = ROOT / "data" / "seed" / "lieux.json"
SHEET = "Restaurants (OSM)"
API = "https://overpass-api.de/api/interpreter"
USER_AGENT = "GoseeContentBot/1.0 (https://gosee-web.vercel.app; curation de contenu)"

# Zone de recherche : le pays du catalogue (code ISO) — rien de figé sur le
# Bénin si le catalogue s'étend à d'autres pays.
PAYS_ISO = "BJ"
RAYON_KM = 3.0
MAX_PAR_LIEU = 3

HEADERS = ["Département", "Catégorie", "Nom", "Description", "Localisation", "GPS (lat, lng)",
           "Lien Google Maps", "Statut", "Horaires", "Téléphone", "Site web", "Note", "Nb avis",
           "Remarque", "Prix", "Gamme de prix", "Durée", "Ambiances", "Idéal pour", "Infos estimées",
           "Alerte", "Photo", "Photo crédit", "Photo source", "OSM id"]

CUISINE_FR = {
    "african": "africaine", "regional": "locale", "local": "locale", "beninese": "béninoise",
    "french": "française", "italian": "italienne", "pizza": "pizzas", "chinese": "chinoise",
    "lebanese": "libanaise", "indian": "indienne", "burger": "burgers", "chicken": "poulet",
    "seafood": "fruits de mer", "fish": "poisson", "grill": "grillades", "barbecue": "grillades",
    "international": "internationale", "vietnamese": "vietnamienne", "japanese": "japonaise",
    "sushi": "sushis", "american": "américaine", "togolese": "togolaise", "nigerian": "nigériane",
    "ivorian": "ivoirienne", "senegalese": "sénégalaise", "coffee_shop": "café", "crepe": "crêpes",
}


def distance_km(a, b):
    r = 6371
    p1, p2 = math.radians(a[0]), math.radians(b[0])
    dp, dl = math.radians(b[0] - a[0]), math.radians(b[1] - a[1])
    h = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(h))


def overpass(query):
    data = urllib.parse.urlencode({"data": query}).encode()
    req = urllib.request.Request(API, data=data, headers={"User-Agent": USER_AGENT})
    for attempt in range(5):
        try:
            with urllib.request.urlopen(req, timeout=180) as resp:
                return json.load(resp)
        except urllib.error.HTTPError as exc:
            if exc.code not in (429, 504) or attempt == 4:
                raise
            time.sleep(15 * (attempt + 1))


def restaurants_osm():
    query = f"""
    [out:json][timeout:170];
    area["ISO3166-1"="{PAYS_ISO}"][admin_level=2]->.pays;
    (
      node["amenity"="restaurant"]["name"](area.pays);
      way["amenity"="restaurant"]["name"](area.pays);
    );
    out center tags;
    """
    out = []
    for el in overpass(query).get("elements", []):
        lat = el.get("lat") or el.get("center", {}).get("lat")
        lon = el.get("lon") or el.get("center", {}).get("lon")
        if lat is None or lon is None:
            continue
        out.append({"id": f"{el['type']}/{el['id']}", "lat": lat, "lng": lon, "tags": el.get("tags", {})})
    return out


def completude(tags):
    return sum(1 for k in ("opening_hours", "phone", "contact:phone", "cuisine", "website", "contact:website", "addr:street") if tags.get(k))


def cuisine_fr(tags):
    raw = [c.strip() for c in re.split(r"[;,]", tags.get("cuisine", "")) if c.strip()]
    fr = [CUISINE_FR.get(c.lower()) for c in raw]
    return [c for c in fr if c]


def description(nom, tags, voisin, d):
    cuisines = cuisine_fr(tags)
    type_ = "Restaurant" if not cuisines else f"Restaurant de cuisine {', '.join(dict.fromkeys(cuisines))}"
    if cuisines == ["café"]:
        type_ = "Café"
    distance = f"{round(d * 1000 / 50) * 50:.0f} m" if d < 1 else f"{d:.1f} km".replace(".", ",")
    return f"{type_} à {distance} de {voisin['nom']} : une adresse pour une pause repas pendant la visite."


def main():
    lieux = json.loads(LIEUX_JSON.read_text(encoding="utf-8"))
    ancres = [l for l in lieux if l["category"] not in ("hotel", "restaurant") and l.get("lat") and not l.get("alerte")]
    print(f"Recherche des restaurants OSM ({PAYS_ISO})...")
    restos = restaurants_osm()
    print(f"{len(restos)} restaurants nommés dans OSM")

    retenus = {}
    for a in ancres:
        proches = []
        for r in restos:
            d = distance_km((a["lat"], a["lng"]), (r["lat"], r["lng"]))
            if d <= RAYON_KM:
                proches.append((completude(r["tags"]), -d, r, d))
        proches.sort(key=lambda x: (x[0], x[1]), reverse=True)
        for _, _, r, d in proches[:MAX_PAR_LIEU]:
            if r["id"] not in retenus or d < retenus[r["id"]][2]:
                retenus[r["id"]] = (r, a, d)

    wb = openpyxl.load_workbook(XLSX_PATH)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    ws.append(HEADERS)
    for r, voisin, d in sorted(retenus.values(), key=lambda x: (x[1]["department"] or "", x[0]["tags"]["name"])):
        t = r["tags"]
        ville = t.get("addr:city") or (voisin["address"] or "").split(",")[0].strip() or voisin["department"]
        rue = t.get("addr:street")
        gps = f"{r['lat']:.7f}, {r['lng']:.7f}"
        ws.append([
            voisin["department"], "Restaurant", t["name"], description(t["name"], t, voisin, d),
            f"{rue}, {ville}" if rue else ville, gps,
            f"https://www.google.com/maps/search/?api=1&query={r['lat']:.7f},{r['lng']:.7f}",
            "OpenStreetMap", t.get("opening_hours"), t.get("phone") or t.get("contact:phone"),
            t.get("website") or t.get("contact:website"), None, None, None,
            None, None, "courte", None, None, "oui", None, None, None, None, r["id"],
        ])
    wb.save(XLSX_PATH)
    print(f"{len(retenus)} restaurants écrits dans la feuille « {SHEET} »")


if __name__ == "__main__":
    main()
