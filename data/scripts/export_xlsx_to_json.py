"""
Exporte benin_contenu_curation.xlsx vers data/seed/lieux.json.

One-off script, hors runtime de l'app Next.js. Nécessite openpyxl
(`pip install openpyxl`).

Usage: python data/scripts/export_xlsx_to_json.py
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX_PATH = ROOT / "benin_contenu_curation.xlsx"
OUT_PATH = ROOT / "data" / "seed" / "lieux.json"

TYPE_TO_CATEGORY = {
    "Site touristique": "site_touristique",
    "Loisir": "loisir",
    "Hôtel": "hotel",
    "Activité": "activite",
    "Événement": "evenement",
}


def parse_gps(raw):
    if not raw:
        return None, None
    match = re.match(r"\s*(-?\d+(?:\.\d+)?)\s*,\s*(-?\d+(?:\.\d+)?)\s*", str(raw))
    if not match:
        return None, None
    return float(match.group(1)), float(match.group(2))


def parse_verified(statut):
    return bool(statut) and "vérifié" in str(statut).lower()


_ACCENTS = str.maketrans(
    "àâäáãåèéêëìíîïòóôöõùúûüçñ",
    "aaaaaaeeeeiiiiooooouuuucn",
)


def slugify(text):
    text = str(text).lower().translate(_ACCENTS)
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text


def assign_unique_slugs(lieux):
    """Slug de base = nom. En cas de collision, suffixe par le département,
    puis par un compteur si ça ne suffit toujours pas."""
    counts = {}
    for lieu in lieux:
        counts[slugify(lieu["nom"])] = counts.get(slugify(lieu["nom"]), 0) + 1

    seen = set()
    for lieu in lieux:
        base = slugify(lieu["nom"])
        slug = base if counts[base] == 1 else f"{base}-{slugify(lieu['department'] or '')}"
        if slug in seen:
            n = 2
            while f"{slug}-{n}" in seen:
                n += 1
            slug = f"{slug}-{n}"
        seen.add(slug)
        lieu["slug"] = slug
    return lieux


def extra_fields(row, headers, offset):
    """Champs d'enrichissement (Horaires, Téléphone, Site web, Note, Nb avis,
    Remarque) ajoutés après les colonnes de base — optionnels, absents pour
    beaucoup de lieux (site naturel/public sans horaires officiels, etc.)."""
    extra = dict(zip(headers[offset:], row[offset:]))
    remarque = extra.get("Remarque")
    description = None
    note = extra.get("Note")
    avis = extra.get("Nb avis")
    return {
        "horaires": extra.get("Horaires"),
        "contact": extra.get("Téléphone"),
        "website": extra.get("Site web"),
        "rating": float(note) if note is not None else None,
        "reviews_count": int(avis) if avis is not None else None,
        "remarque": remarque,
    }


def rows_from_sheet_1(ws):
    """60 Lieux incontournables: Département, #, Nom, Type, Description,
    Localisation, GPS, Lien Google Maps, Statut, [+ enrichissement]."""
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    for row in ws.iter_rows(min_row=2, values_only=True):
        department, _num, nom, type_, description, localisation, gps, maps_url, statut = row[:9]
        if not nom:
            continue
        lat, lng = parse_gps(gps)
        yield {
            "nom": nom,
            "category": TYPE_TO_CATEGORY.get(type_, "activite"),
            "description": description,
            "department": department,
            "country": "Bénin",
            "address": localisation,
            "lat": lat,
            "lng": lng,
            "google_maps_url": maps_url,
            "verified": parse_verified(statut),
            "source": statut,
            **extra_fields(row, headers, 9),
        }


def rows_from_sheet_2(ws):
    """Hôtels-Activités-Événements: Département, Catégorie, Nom, Description,
    Localisation, GPS, Lien Google Maps, Statut, [+ enrichissement]."""
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    for row in ws.iter_rows(min_row=2, values_only=True):
        department, categorie, nom, description, localisation, gps, maps_url, statut = row[:8]
        if not nom:
            continue
        lat, lng = parse_gps(gps)
        yield {
            "nom": nom,
            "category": TYPE_TO_CATEGORY.get(categorie, "activite"),
            "description": description,
            "department": department,
            "country": "Bénin",
            "address": localisation,
            "lat": lat,
            "lng": lng,
            "google_maps_url": maps_url,
            "verified": parse_verified(statut),
            "source": statut,
            **extra_fields(row, headers, 8),
        }


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    lieux = list(rows_from_sheet_1(wb.worksheets[0])) + list(rows_from_sheet_2(wb.worksheets[1]))

    # La colonne "Remarque" (fermeture temporaire, statut incertain...) n'a
    # pas sa propre colonne en base : elle rejoint la description, visible
    # par l'utilisateur, plutôt que d'être silencieusement perdue.
    for lieu in lieux:
        remarque = lieu.pop("remarque", None)
        if remarque:
            lieu["description"] = f"{lieu['description']} ({remarque})" if lieu["description"] else remarque

    seen = set()
    for lieu in lieux:
        key = (lieu["nom"], lieu["department"])
        if key in seen:
            print(f"Doublon détecté: {key}")
        seen.add(key)

    assign_unique_slugs(lieux)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(json.dumps(lieux, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{len(lieux)} lieux exportés vers {OUT_PATH}")


if __name__ == "__main__":
    main()
