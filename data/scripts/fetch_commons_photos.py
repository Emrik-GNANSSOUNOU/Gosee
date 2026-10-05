"""
Cherche de vraies photos des lieux sur Wikimedia Commons (photos libres de
droits, sans clé API ni facturation) et les propose dans
benin_contenu_curation.xlsx (colonnes « Photo », « Photo crédit »,
« Photo source »).

Deux recherches par lieu :
  1. géographique : fichiers photographiés à proximité de ses coordonnées GPS ;
  2. textuelle : fichiers dont le titre ou la description cite son nom.
Les candidats sont notés (nom cité dans le titre, proximité, taille, format
paysage) et filtrés par licence (CC0, domaine public, CC BY, CC BY-SA) :
seules les licences qui autorisent la réutilisation commerciale avec
attribution sont retenues, et le crédit exigé par la licence est enregistré.

Par défaut, le script N'ÉCRIT RIEN dans le classeur : il produit
data/photos_commons_revue.csv pour vérifier à l'œil chaque photo proposée
(lien vers la page Commons). Avec --apply, il remplit les colonnes Photo des
lieux qui n'en ont pas encore (une photo saisie à la main n'est jamais
écrasée), en ne gardant que les candidats au-dessus du seuil de confiance.

Réseau requis : commons.wikimedia.org (API) — à autoriser dans la politique
réseau de l'environnement si besoin. Respecte la politique d'usage de
Wikimedia (User-Agent identifiant, requêtes séquentielles, pause entre deux).

Usage :
  python data/scripts/fetch_commons_photos.py            # revue seule
  python data/scripts/fetch_commons_photos.py --apply    # + écrit le xlsx
puis : python data/scripts/export_xlsx_to_json.py
       python data/scripts/generate_sql_update.py
"""
import argparse
import csv
import html
import json
import math
import re
import sys
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX_PATH = ROOT / "benin_contenu_curation.xlsx"
REVIEW_PATH = ROOT / "data" / "photos_commons_revue.csv"
API = "https://commons.wikimedia.org/w/api.php"
USER_AGENT = "GoseeContentBot/1.0 (https://gosee-web.vercel.app; curation de contenu)"

ALLOWED_LICENSE = re.compile(r"^(cc0|public domain|pd\b|cc by(-sa)? \d)", re.I)
EXCLUDED_TITLE = re.compile(r"\b(map|carte|logo|flag|drapeau|coat of arms|blason|locator|plan)\b", re.I)
RADIUS_M = {"Site touristique": 1500, "Loisir": 3000, "Activité": 1500, "Hôtel": 300}
MIN_SCORE = 4.0
COLUMNS = ["Photo", "Photo crédit", "Photo source"]

STOPWORDS = {"de", "du", "des", "la", "le", "les", "l", "d", "a", "au", "aux", "et", "en", "sur",
             "visite", "guidee", "sortie", "balade", "decouverte", "hotel", "parc", "national",
             "village", "ville", "musee", "foret", "classee", "the", "of"}


def normalize(text):
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode().lower()
    return re.findall(r"[a-z0-9]+", text)


def name_tokens(nom, localisation=""):
    # Les mots de la ville/du département ne comptent pas : sinon « Rio Hôtel
    # Ouidah » reconnaîtrait n'importe quelle photo titrée « … Ouidah ».
    lieu = set(normalize(localisation or ""))
    return {t for t in normalize(nom) if len(t) > 2 and t not in STOPWORDS and t not in lieu}


def api(params):
    params = {**params, "format": "json", "formatversion": "2"}
    url = f"{API}?{urllib.parse.urlencode(params)}"
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    # Wikimedia limite le débit (429) : on attend le délai qu'il indique
    # (Retry-After) ou un délai croissant, puis on réessaie.
    for attempt in range(6):
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                data = json.load(resp)
            time.sleep(1.5)
            return data
        except urllib.error.HTTPError as exc:
            if exc.code != 429 or attempt == 5:
                raise
            wait = int(exc.headers.get("Retry-After") or 0) or 5 * 2 ** attempt
            time.sleep(min(wait, 120))


IMAGEINFO = {
    "prop": "imageinfo|coordinates",
    "iiprop": "url|size|mime|extmetadata",
    "iiurlwidth": "1600",
    "iiextmetadatafilter": "Artist|LicenseShortName|LicenseUrl|ImageDescription",
}


def candidates_near(lat, lng, radius):
    data = api({"action": "query", "generator": "geosearch", "ggscoord": f"{lat}|{lng}",
                "ggsradius": str(radius), "ggsnamespace": "6", "ggslimit": "30", **IMAGEINFO})
    return data.get("query", {}).get("pages", [])


def candidates_named(nom, ville):
    data = api({"action": "query", "generator": "search", "gsrsearch": f'"{nom}" {ville} filetype:bitmap',
                "gsrnamespace": "6", "gsrlimit": "20", **IMAGEINFO})
    return data.get("query", {}).get("pages", [])


def strip_html(text):
    return html.unescape(re.sub(r"<[^>]+>", "", text or "")).strip()


def distance_m(lat1, lng1, lat2, lng2):
    r = 6371000
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = math.radians(lat2 - lat1), math.radians(lng2 - lng1)
    h = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(h))


def score(page, tokens, lat, lng, radius):
    info = (page.get("imageinfo") or [{}])[0]
    meta = info.get("extmetadata", {})
    title = page.get("title", "")
    license_name = meta.get("LicenseShortName", {}).get("value", "")
    if info.get("mime") != "image/jpeg" or EXCLUDED_TITLE.search(title):
        return None
    if not ALLOWED_LICENSE.search(license_name):
        return None
    width, height = info.get("width", 0), info.get("height", 0)
    if width < 1000 or height < 600:
        return None

    text_tokens = set(normalize(title)) | set(normalize(strip_html(meta.get("ImageDescription", {}).get("value", ""))))
    matched = len(tokens & text_tokens)
    s = 3.0 * matched
    if width >= height:
        s += 1.0  # les visuels de l'app sont en paysage
    s += min(width, 4000) / 4000
    coords = page.get("coordinates") or []
    if coords and lat is not None:
        d = distance_m(lat, lng, coords[0]["lat"], coords[0]["lon"])
        s += 2.0 * max(0.0, 1 - d / radius)
    return {
        "score": round(s, 2),
        "matched": matched,
        # Sans les paramètres de suivi (?utm_source=...) ajoutés par l'API.
        "url": (info.get("thumburl") or info.get("url") or "").split("?")[0],
        "source": info.get("descriptionurl"),
        "credit": f"Photo : {strip_html(meta.get('Artist', {}).get('value', '')) or 'auteur inconnu'} — {license_name}, via Wikimedia Commons",
        "title": title,
    }


def best_photo(nom, type_, localisation, lat, lng):
    tokens = name_tokens(nom, localisation)
    radius = RADIUS_M.get(type_, 1500)
    ville = (localisation or "").split(",")[0]
    pages = {}
    if lat is not None:
        for p in candidates_near(lat, lng, radius):
            pages[p["title"]] = p
    for p in candidates_named(nom, ville):
        pages.setdefault(p["title"], p)

    scored = [s for p in pages.values() if (s := score(p, tokens, lat, lng, radius))]
    # Le nom du lieu doit être cité dans le titre ou la description du
    # fichier : la seule proximité GPS ramène trop de photos d'autre chose.
    scored = [s for s in scored if s["matched"] > 0]
    scored.sort(key=lambda s: s["score"], reverse=True)
    return scored[0] if scored else None


def parse_gps(raw):
    m = re.match(r"\s*(-?\d+(?:\.\d+)?)\s*,\s*(-?\d+(?:\.\d+)?)", str(raw or ""))
    return (float(m.group(1)), float(m.group(2))) if m else (None, None)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--apply", action="store_true", help="écrire les photos retenues dans le xlsx")
    parser.add_argument("--min-score", type=float, default=MIN_SCORE)
    args = parser.parse_args()

    wb = openpyxl.load_workbook(XLSX_PATH)
    rows_out = []
    written = 0
    for ws in wb.worksheets:
        headers = [c.value for c in ws[1]]
        for col in COLUMNS:
            if col not in headers:
                ws.cell(1, len(headers) + 1).value = col
                headers.append(col)
        idx = {h: i for i, h in enumerate(headers)}
        type_col = "Type" if "Type" in idx else "Catégorie"

        for row in ws.iter_rows(min_row=2):
            nom = row[idx["Nom"]].value
            if not nom:
                continue
            type_ = row[idx[type_col]].value
            lat, lng = parse_gps(row[idx["GPS (lat, lng)"]].value)
            try:
                photo = best_photo(nom, type_, row[idx["Localisation"]].value, lat, lng)
            except Exception as exc:  # réseau, quota... : on continue
                print(f"  ! {nom} : {exc}", file=sys.stderr)
                photo = None
            retenu = bool(photo and photo["score"] >= args.min_score)
            deja = bool(row[idx["Photo"]].value)
            rows_out.append({
                "nom": nom, "retenu": "oui" if retenu else "non", "deja_une_photo": "oui" if deja else "non",
                "score": photo["score"] if photo else "", "fichier": photo["title"] if photo else "",
                "page_commons": photo["source"] if photo else "", "credit": photo["credit"] if photo else "",
            })
            print(f"{'✓' if retenu else '·'} {nom} -> {photo['title'] if photo else 'aucune'}")
            if args.apply and retenu and not deja:
                row[idx["Photo"]].value = photo["url"]
                row[idx["Photo crédit"]].value = photo["credit"]
                row[idx["Photo source"]].value = photo["source"]
                written += 1

    REVIEW_PATH.parent.mkdir(parents=True, exist_ok=True)
    with REVIEW_PATH.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows_out[0].keys()))
        writer.writeheader()
        writer.writerows(rows_out)
    print(f"\nRevue : {REVIEW_PATH}")
    if args.apply:
        wb.save(XLSX_PATH)
        print(f"{written} photos écrites dans {XLSX_PATH.name}")


if __name__ == "__main__":
    main()
